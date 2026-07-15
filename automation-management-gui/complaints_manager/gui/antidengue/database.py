from __future__ import annotations

import csv
import hashlib
import json
import re
import shlex
import shutil
import sqlite3
import subprocess
import tempfile
import threading
import time
import urllib.request
import uuid
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path

from imgui_bundle import imgui
from openpyxl import load_workbook

from complaints_manager.core.commands import (
    ArgField,
    COMMANDS,
    SYSTEMS,
    CommandSpec,
    SystemSpec,
    command_by_id,
    commands_by_system,
)
from complaints_manager.core.services import normalize_whatsapp_account_id
from complaints_manager.gui.paths import (
    apply_selected_folder,
    apply_selected_file,
    consume_file_dialog_result,
    consume_folder_dialog_result,
    existing_initial_dir,
    resolve_gui_path,
    start_file_dialog,
    start_folder_dialog,
)
from complaints_manager.gui.state import GuiState

DISPATCH_SETTING_ID = "b50691f465c2e1a"

MASTER_SCHOOL_SHIFTS = ("Single", "Morning", "Evening", "Double")

MASTER_SCHOOL_TYPES = ("Male", "Female", "Co-Education")

MASTER_SCHOOL_LEVELS = ("Primary", "Middle", "High", "Higher Secondary", "sMosque")

from complaints_manager.gui.antidengue.health_repository import open_pocketbase_read_connection
from complaints_manager.gui.antidengue.runtime import antidengue_pocketbase_root
from complaints_manager.gui.common.formatting import _option_index, _safe_json_loads

def _pocketbase_report_runs_has_lifecycle(conn: sqlite3.Connection) -> bool:
    columns = {
        str(row[1])
        for row in conn.execute("pragma table_info(report_runs)").fetchall()
    }
    return {"archived", "archived_at"}.issubset(columns)

def _now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")

def antidengue_record_id(prefix: str, *parts: object) -> str:
    key = ":".join(str(part) for part in parts if part is not None)
    return hashlib.sha1(f"{prefix}:{key}".encode("utf-8")).hexdigest()[:15]

@contextmanager
def open_pocketbase_write_connection(db_path: Path):
    conn = sqlite3.connect(str(db_path), timeout=5)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

def _pb_bool(value: object, default: bool = False) -> bool:
    if value is None:
        return default
    text = str(value).strip().lower()
    if not text:
        return default
    return text not in {"0", "false", "no", "off"}

def _one_or_zero(value: object) -> int:
    return 1 if _pb_bool(value) else 0

def _dispatch_settings_defaults() -> dict[str, object]:
    return {
        "id": DISPATCH_SETTING_ID,
        "active": True,
        "name": "default",
        "allow_individual": True,
        "allow_groups": True,
        "manual_only": False,
        "attach_excel": True,
        "send_failed_only": False,
        "require_preview": False,
        "notes": "Default AntiDengue dispatch controls.",
    }

def load_antidengue_dispatch_settings(state: GuiState) -> tuple[dict[str, object], str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    defaults = _dispatch_settings_defaults()
    if not db_path.is_file():
        return defaults, f"PocketBase database has not been initialized yet: {db_path}"
    try:
        with open_pocketbase_read_connection(db_path, state) as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    "select name from sqlite_master where type = 'table'"
                ).fetchall()
            }
            if "dispatch_settings" not in tables:
                return defaults, "Dispatch Center schema is missing. Apply PocketBase migrations."
            row = conn.execute(
                """
                select id, active, name, allow_individual, allow_groups, manual_only,
                       attach_excel, send_failed_only, require_preview, notes
                from dispatch_settings
                where active = 1
                order by name
                limit 1
                """
            ).fetchone()
            if row is None:
                return defaults, ""
            settings = dict(defaults)
            settings.update(dict(row))
            for key in (
                "active",
                "allow_individual",
                "allow_groups",
                "manual_only",
                "attach_excel",
                "send_failed_only",
                "require_preview",
            ):
                settings[key] = _pb_bool(settings.get(key), bool(defaults[key]))
            return settings, ""
    except sqlite3.Error as exc:
        return defaults, f"Could not read dispatch settings: {exc}"

def save_antidengue_dispatch_settings(
    state: GuiState,
    settings: dict[str, object],
) -> tuple[bool, str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    try:
        with open_pocketbase_write_connection(db_path) as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    "select name from sqlite_master where type = 'table'"
                ).fetchall()
            }
            if "dispatch_settings" not in tables:
                return False, "Dispatch Center schema is missing. Apply PocketBase migrations."
            values = _dispatch_settings_defaults()
            values.update(settings)
            conn.execute(
                """
                insert into dispatch_settings
                (id, active, name, allow_individual, allow_groups, manual_only,
                 attach_excel, send_failed_only, require_preview, notes)
                values
                (:id, :active, :name, :allow_individual, :allow_groups, :manual_only,
                 :attach_excel, :send_failed_only, :require_preview, :notes)
                on conflict(name) do update set
                  active = excluded.active,
                  allow_individual = excluded.allow_individual,
                  allow_groups = excluded.allow_groups,
                  manual_only = excluded.manual_only,
                  attach_excel = excluded.attach_excel,
                  send_failed_only = excluded.send_failed_only,
                  require_preview = excluded.require_preview,
                  notes = excluded.notes
                """,
                {
                    "id": str(values["id"]),
                    "active": _one_or_zero(values["active"]),
                    "name": str(values["name"] or "default"),
                    "allow_individual": _one_or_zero(values["allow_individual"]),
                    "allow_groups": _one_or_zero(values["allow_groups"]),
                    "manual_only": _one_or_zero(values["manual_only"]),
                    "attach_excel": _one_or_zero(values["attach_excel"]),
                    "send_failed_only": _one_or_zero(values["send_failed_only"]),
                    "require_preview": _one_or_zero(values["require_preview"]),
                    "notes": str(values.get("notes") or ""),
                },
            )
        return True, "Saved AntiDengue dispatch settings."
    except sqlite3.Error as exc:
        return False, f"Could not save dispatch settings: {exc}"

def load_antidengue_dispatch_groups(state: GuiState) -> tuple[list[dict], str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    if not db_path.is_file():
        return [], f"PocketBase database has not been initialized yet: {db_path}"
    try:
        with open_pocketbase_read_connection(db_path, state) as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    "select name from sqlite_master where type = 'table'"
                ).fetchall()
            }
            if "whatsapp_groups" not in tables:
                return [], "Dispatch group schema is missing. Apply PocketBase migrations."
            rows = conn.execute(
                """
                select
                  wg.id,
                  wg.enabled,
                  wg.name,
                  wg.target,
                  coalesce(wg.route_kind, 'district') as route_kind,
                  coalesce(wg.message_mode, 'full_report') as message_mode,
                  coalesce(wg.attach_excel, 1) as attach_excel,
                  coalesce(wg.manual_only, 0) as manual_only,
                  coalesce(wg.attachment_text_mode, '') as attachment_text_mode,
                  coalesce(wg.delay_ms, 0) as delay_ms,
                  coalesce(wg.tehsil_ref, '') as tehsil_ref,
                  coalesce(wg.markaz_ref, '') as markaz_ref,
                  coalesce(t.name, '') as tehsil,
                  coalesce(m.name, '') as markaz,
                  coalesce(wg.notes, '') as notes
                from whatsapp_groups wg
                left join tehsils t on t.id = wg.tehsil_ref
                left join markazes m on m.id = wg.markaz_ref
                order by wg.enabled desc, wg.route_kind, wg.name
                """
            ).fetchall()
            return [dict(row) for row in rows], ""
    except sqlite3.Error as exc:
        return [], f"Could not read dispatch groups: {exc}"

def load_antidengue_scope_options(state: GuiState) -> tuple[list[dict], list[dict], str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    try:
        with open_pocketbase_read_connection(db_path, state) as conn:
            tehsils = [
                dict(row)
                for row in conn.execute(
                    "select id, name from tehsils where active = 1 order by name"
                ).fetchall()
            ]
            markazes = [
                dict(row)
                for row in conn.execute(
                    """
                    select m.id, m.name, coalesce(m.tehsil_ref, '') as tehsil_ref, coalesce(t.name, '') as tehsil
                    from markazes m
                    left join tehsils t on t.id = m.tehsil_ref
                    where m.active = 1
                    order by t.name, m.name
                    """
                ).fetchall()
            ]
            return tehsils, markazes, ""
    except sqlite3.Error as exc:
        return [], [], f"Could not read scope options: {exc}"

def _normalized_emis(value: object) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())

def normalize_pk_mobile_for_gui(value: object) -> str:
    digits = "".join(ch for ch in str(value or "") if ch.isdigit())
    if not digits:
        return ""
    if digits.startswith("0092"):
        digits = digits[2:]
    if digits.startswith("92"):
        return digits
    if digits.startswith("0"):
        return "92" + digits[1:]
    if len(digits) == 10 and digits.startswith("3"):
        return "92" + digits
    return digits

def _school_level_from_name_gui(name: object) -> str:
    prefix = str(name or "").strip().upper().split(" ", 1)[0]
    if prefix == "GPS":
        return "Primary"
    if prefix == "GES":
        return "Middle"
    if prefix == "GMMS":
        return "sMosque"
    return ""

def _pocketbase_table_columns(conn: sqlite3.Connection, table: str) -> set[str]:
    return {str(row[1]) for row in conn.execute(f"pragma table_info({table})").fetchall()}

def _master_school_form_defaults() -> dict[str, object]:
    return {
        "emis": "",
        "name": "",
        "district_ref": "",
        "department_ref": "",
        "wing_ref": "",
        "tehsil_ref": "",
        "markaz_ref": "",
        "shift": "Single",
        "school_type": "Male",
        "deos_wise": "M-EE",
        "school_level": "Primary",
        "head_name": "",
        "head_contact": "",
        "active": True,
        "notes": "",
        "last_result": "",
        "last_success": False,
        "last_saved_emis": "",
    }

def antidengue_new_school_form(state: GuiState) -> dict[str, object]:
    form = state.app_settings.setdefault("antidengue_new_school_form", {})
    if not isinstance(form, dict):
        form = {}
        state.app_settings["antidengue_new_school_form"] = form
    for key, value in _master_school_form_defaults().items():
        form.setdefault(key, value)
    return form

def load_antidengue_master_data_options(state: GuiState) -> tuple[dict[str, list[dict]], str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    if not db_path.is_file():
        return {}, f"PocketBase database has not been initialized yet: {db_path}"
    try:
        with open_pocketbase_read_connection(db_path, state) as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    "select name from sqlite_master where type = 'table'"
                ).fetchall()
            }
            required = {"districts", "departments", "wings", "tehsils", "markazes", "schools"}
            missing = sorted(required.difference(tables))
            if missing:
                return {}, "PocketBase is missing master collections: " + ", ".join(missing)

            options = {
                "districts": [
                    dict(row)
                    for row in conn.execute(
                        "select id, name from districts where active = 1 order by name"
                    ).fetchall()
                ],
                "departments": [
                    dict(row)
                    for row in conn.execute(
                        "select id, name from departments where active = 1 order by name"
                    ).fetchall()
                ],
                "wings": [
                    dict(row)
                    for row in conn.execute(
                        """
                        select id, name, district_ref, department_ref
                        from wings
                        where active = 1
                        order by name
                        """
                    ).fetchall()
                ],
                "tehsils": [
                    dict(row)
                    for row in conn.execute(
                        """
                        select id, name, district_ref
                        from tehsils
                        where active = 1
                        order by name
                        """
                    ).fetchall()
                ],
                "markazes": [
                    dict(row)
                    for row in conn.execute(
                        """
                        select m.id, m.name, m.tehsil_ref, m.wing_ref, coalesce(t.name, '') as tehsil
                        from markazes m
                        left join tehsils t on t.id = m.tehsil_ref
                        where m.active = 1
                        order by t.name, m.name
                        """
                    ).fetchall()
                ],
            }
            return options, ""
    except sqlite3.Error as exc:
        return {}, f"Could not read master data options: {exc}"

def _relation_options(rows: list[dict], placeholder: str, *, include_parent: bool = False) -> list[tuple[str, str]]:
    options = [("", placeholder)]
    for row in rows:
        label = str(row.get("name") or "-")
        if include_parent and row.get("tehsil"):
            label = f"{row.get('tehsil')} / {label}"
        options.append((str(row.get("id") or ""), label))
    return options

def _set_single_default(form: dict[str, object], key: str, rows: list[dict]) -> None:
    if not form.get(key) and len(rows) == 1:
        form[key] = str(rows[0].get("id") or "")

def _draw_relation_combo(
    label: str,
    key: str,
    form: dict[str, object],
    options: list[tuple[str, str]],
    *,
    width: int,
) -> bool:
    valid_values = {value for value, _ in options}
    current = str(form.get(key) or "")
    if current not in valid_values:
        current = ""
        form[key] = ""
    index = _option_index(options, current)
    imgui.set_next_item_width(width)
    changed, index = imgui.combo(label, index, [option_label for _, option_label in options])
    if changed:
        form[key] = options[index][0]
    return changed

def resolve_antidengue_school_owners(
    state: GuiState,
    *,
    district_ref: str,
    department_ref: str,
    wing_ref: str,
    tehsil_ref: str,
    markaz_ref: str,
) -> tuple[dict[str, list[dict]], str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    try:
        with open_pocketbase_read_connection(db_path, state) as conn:
            ddeo_rows = [
                dict(row)
                for row in conn.execute(
                    """
                    select distinct
                      d.id,
                      d.name,
                      d.normalized_mobile,
                      j.id as jurisdiction_id
                    from ddeo_jurisdictions j
                    join ddeo_officers d on d.id = j.ddeo_ref and d.active = 1
                    where j.active = 1
                      and j.district_ref = ?
                      and j.department_ref = ?
                      and j.wing_ref = ?
                      and j.tehsil_ref = ?
                    order by d.name, d.normalized_mobile
                    """,
                    (district_ref, department_ref, wing_ref, tehsil_ref),
                ).fetchall()
            ]
            aeo_rows = [
                dict(row)
                for row in conn.execute(
                    """
                    select distinct
                      a.id,
                      a.name,
                      a.normalized_mobile,
                      j.id as jurisdiction_id
                    from aeo_jurisdictions j
                    join aeo_officers a on a.id = j.aeo_ref and a.active = 1
                    where j.active = 1
                      and j.district_ref = ?
                      and j.department_ref = ?
                      and j.wing_ref = ?
                      and j.markaz_ref = ?
                      and (j.tehsil_ref = '' or j.tehsil_ref = ?)
                    order by a.name, a.normalized_mobile
                    """,
                    (district_ref, department_ref, wing_ref, markaz_ref, tehsil_ref),
                ).fetchall()
            ]
            return {"ddeo": ddeo_rows, "aeo": aeo_rows}, ""
    except sqlite3.Error as exc:
        return {"ddeo": [], "aeo": []}, f"Could not resolve school officers: {exc}"

def validate_antidengue_new_school_form(
    state: GuiState,
    form: dict[str, object],
) -> tuple[list[str], dict[str, list[dict]], str]:
    errors: list[str] = []
    emis = _normalized_emis(form.get("emis"))
    if not emis:
        errors.append("EMIS is required.")
    elif len(emis) != 8:
        errors.append("EMIS must be 8 digits.")
    if not str(form.get("name") or "").strip():
        errors.append("School name is required.")

    required_refs = {
        "district_ref": "District",
        "department_ref": "Department",
        "wing_ref": "Wing",
        "tehsil_ref": "Tehsil",
        "markaz_ref": "Markaz",
    }
    for key, label in required_refs.items():
        if not str(form.get(key) or "").strip():
            errors.append(f"{label} is required.")

    if not str(form.get("shift") or "").strip():
        errors.append("Shift is required.")
    if not str(form.get("school_type") or "").strip():
        errors.append("School type is required.")
    if not str(form.get("school_level") or "").strip():
        errors.append("School level is required.")
    if not str(form.get("head_name") or "").strip():
        errors.append("Head name is required.")

    head_contact = normalize_pk_mobile_for_gui(form.get("head_contact"))
    if not head_contact:
        errors.append("Head contact is required.")
    elif len(head_contact) != 12 or not head_contact.startswith("923"):
        errors.append("Head contact must normalize to a Pakistani WhatsApp number, e.g. 923001234567.")

    owners = {"ddeo": [], "aeo": []}
    owner_error = ""
    refs_complete = all(str(form.get(key) or "").strip() for key in required_refs)
    if refs_complete:
        owners, owner_error = resolve_antidengue_school_owners(
            state,
            district_ref=str(form.get("district_ref") or ""),
            department_ref=str(form.get("department_ref") or ""),
            wing_ref=str(form.get("wing_ref") or ""),
            tehsil_ref=str(form.get("tehsil_ref") or ""),
            markaz_ref=str(form.get("markaz_ref") or ""),
        )
        if owner_error:
            errors.append(owner_error)
        if len(owners.get("ddeo") or []) != 1:
            errors.append("Selected hierarchy must resolve to exactly one active DDEO.")
        if len(owners.get("aeo") or []) != 1:
            errors.append("Selected hierarchy must resolve to exactly one active AEO.")

    return errors, owners, owner_error

def create_antidengue_school_from_form(
    state: GuiState,
    form: dict[str, object],
) -> tuple[bool, str]:
    errors, _owners, _owner_error = validate_antidengue_new_school_form(state, form)
    if errors:
        return False, "Cannot save school: " + " ".join(errors[:3])

    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    emis = _normalized_emis(form.get("emis"))
    school_name = " ".join(str(form.get("name") or "").strip().upper().split())
    head_contact = normalize_pk_mobile_for_gui(form.get("head_contact"))
    source_hash_payload = {
        "emis": emis,
        "name": school_name,
        "district_ref": str(form.get("district_ref") or ""),
        "department_ref": str(form.get("department_ref") or ""),
        "wing_ref": str(form.get("wing_ref") or ""),
        "tehsil_ref": str(form.get("tehsil_ref") or ""),
        "markaz_ref": str(form.get("markaz_ref") or ""),
        "source": "manual_gui",
    }
    values: dict[str, object] = {
        "id": antidengue_record_id("school", emis),
        "active": _one_or_zero(form.get("active")),
        "emis": emis,
        "name": school_name,
        "shift": str(form.get("shift") or "").strip(),
        "head_name": " ".join(str(form.get("head_name") or "").strip().split()),
        "head_contact": head_contact,
        "school_type": str(form.get("school_type") or "").strip(),
        "deos_wise": str(form.get("deos_wise") or "").strip() or "M-EE",
        "school_level": str(form.get("school_level") or "").strip(),
        "source_row_hash": hashlib.sha256(
            json.dumps(source_hash_payload, sort_keys=True).encode("utf-8")
        ).hexdigest(),
        "district_ref": str(form.get("district_ref") or ""),
        "department_ref": str(form.get("department_ref") or ""),
        "wing_ref": str(form.get("wing_ref") or ""),
        "tehsil_ref": str(form.get("tehsil_ref") or ""),
        "markaz_ref": str(form.get("markaz_ref") or ""),
        "source": "manual_gui",
        "notes": str(form.get("notes") or "").strip(),
    }

    try:
        with open_pocketbase_write_connection(db_path) as conn:
            existing = conn.execute(
                "select id, name, active from schools where emis = ?",
                (emis,),
            ).fetchone()
            if existing:
                status = "active" if _pb_bool(existing["active"]) else "inactive"
                return False, f"EMIS {emis} already exists as {existing['name']} ({status})."

            columns_in_table = _pocketbase_table_columns(conn, "schools")
            filtered = {
                key: value
                for key, value in values.items()
                if key in columns_in_table
            }
            columns = ", ".join(filtered)
            placeholders = ", ".join(f":{key}" for key in filtered)
            conn.execute(
                f"insert into schools ({columns}) values ({placeholders})",
                filtered,
            )
        state.cached_metrics.pop("antidengue", None)
        return True, f"Saved school {emis} - {school_name}."
    except sqlite3.Error as exc:
        return False, f"Could not save school: {exc}"

def update_antidengue_dispatch_group(
    state: GuiState,
    group_id: str,
    updates: dict[str, object],
) -> tuple[bool, str]:
    allowed = {
        "enabled",
        "name",
        "target",
        "route_kind",
        "message_mode",
        "attach_excel",
        "manual_only",
        "attachment_text_mode",
        "delay_ms",
        "tehsil_ref",
        "markaz_ref",
        "notes",
    }
    clean_updates = {key: value for key, value in updates.items() if key in allowed}
    if not clean_updates:
        return True, "No group changes to save."

    for key in ("enabled", "attach_excel", "manual_only"):
        if key in clean_updates:
            clean_updates[key] = _one_or_zero(clean_updates[key])
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    try:
        with open_pocketbase_write_connection(db_path) as conn:
            assignments = ", ".join(f"{key} = :{key}" for key in clean_updates)
            conn.execute(
                f"update whatsapp_groups set {assignments} where id = :id",
                {**clean_updates, "id": group_id},
            )
        return True, "Saved dispatch group."
    except sqlite3.Error as exc:
        return False, f"Could not save dispatch group: {exc}"

def add_antidengue_dispatch_group(
    state: GuiState,
    values: dict[str, object],
) -> tuple[bool, str]:
    target = str(values.get("target") or "").strip()
    name = str(values.get("name") or "").strip()
    if not name:
        return False, "Enter a group name."
    if not target.endswith("@g.us"):
        return False, "Group target must be a WhatsApp group JID ending with @g.us."

    route_kind = str(values.get("route_kind") or "district")
    message_mode = str(values.get("message_mode") or "full_report")
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    try:
        with open_pocketbase_write_connection(db_path) as conn:
            conn.execute(
                """
                insert into whatsapp_groups
                (id, enabled, name, target, route_kind, message_mode, attach_excel,
                 manual_only, attachment_text_mode, delay_ms, tehsil_ref, markaz_ref, notes)
                values
                (:id, :enabled, :name, :target, :route_kind, :message_mode, :attach_excel,
                 :manual_only, :attachment_text_mode, :delay_ms, :tehsil_ref, :markaz_ref, :notes)
                """,
                {
                    "id": antidengue_record_id("whatsapp_group", target),
                    "enabled": _one_or_zero(values.get("enabled", True)),
                    "name": name,
                    "target": target,
                    "route_kind": route_kind,
                    "message_mode": message_mode,
                    "attach_excel": _one_or_zero(values.get("attach_excel", True)),
                    "manual_only": _one_or_zero(values.get("manual_only", False)),
                    "attachment_text_mode": str(
                        values.get("attachment_text_mode") or "separate"
                    ),
                    "delay_ms": int(values.get("delay_ms") or 1500),
                    "tehsil_ref": str(values.get("tehsil_ref") or ""),
                    "markaz_ref": str(values.get("markaz_ref") or ""),
                    "notes": str(values.get("notes") or "Created from Dispatch Center."),
                },
            )
        return True, f"Added dispatch group: {name}."
    except sqlite3.IntegrityError:
        return False, "A dispatch group with this target already exists."
    except (ValueError, sqlite3.Error) as exc:
        return False, f"Could not add dispatch group: {exc}"

def record_antidengue_dispatch_event(
    state: GuiState,
    *,
    run_key: str,
    channel: str,
    action: str,
    target: str,
    recipient_name: str,
    status: str,
    message: str,
    payload: dict,
) -> None:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    try:
        with open_pocketbase_write_connection(db_path) as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    "select name from sqlite_master where type = 'table'"
                ).fetchall()
            }
            if "dispatch_events" not in tables:
                return
            event_id = antidengue_record_id(
                "dispatch_event",
                run_key,
                channel,
                action,
                target,
                time.time_ns(),
            )
            conn.execute(
                """
                insert into dispatch_events
                (id, run_key, channel, action, target, recipient_name, status,
                 message_preview, payload, created_at, notes)
                values
                (:id, :run_key, :channel, :action, :target, :recipient_name, :status,
                 :message_preview, :payload, :created_at, :notes)
                """,
                {
                    "id": event_id,
                    "run_key": run_key,
                    "channel": channel,
                    "action": action,
                    "target": target,
                    "recipient_name": recipient_name,
                    "status": status,
                    "message_preview": message[:1000],
                    "payload": json.dumps(payload, ensure_ascii=False, sort_keys=True),
                    "created_at": _now_iso(),
                    "notes": "Recorded from Automation Management GUI.",
                },
            )
    except sqlite3.Error:
        return

def load_antidengue_latest_dispatch_context(state: GuiState) -> tuple[dict, str]:
    db_path = antidengue_pocketbase_root(state) / "pb_data" / "data.db"
    if not db_path.is_file():
        return {}, f"PocketBase database has not been initialized yet: {db_path}"
    try:
        with open_pocketbase_read_connection(db_path, state) as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    "select name from sqlite_master where type = 'table'"
                ).fetchall()
            }
            required = {"report_runs", "dormant_records", "schools", "tehsils", "markazes"}
            missing = required.difference(tables)
            if missing:
                return {}, "PocketBase is missing collections: " + ", ".join(sorted(missing))
            has_lifecycle = _pocketbase_report_runs_has_lifecycle(conn)
            row = conn.execute(
                f"""
                select started_at, raw_file_name, summary
                from report_runs
                {"where coalesce(archived, 0) = 0" if has_lifecycle else ""}
                order by started_at desc
                limit 1
                """
            ).fetchone()
            if row is None:
                return {}, "No AntiDengue run history is available yet."
            run_key = str(row["started_at"] or "")
            dormant_rows = [
                dict(item)
                for item in conn.execute(
                    """
                    select
                      s.id as school_ref,
                      s.emis,
                      s.name as school_name,
                      s.tehsil_ref,
                      s.markaz_ref,
                      t.name as tehsil,
                      m.name as markaz,
                      coalesce(override_ddeo.name, jurisdiction_ddeo.name, '') as ddeo_name,
                      coalesce(override_ddeo.normalized_mobile, jurisdiction_ddeo.normalized_mobile, '') as ddeo_mobile,
                      coalesce(override_aeo.name, jurisdiction_aeo.name, '') as aeo_name,
                      coalesce(override_aeo.normalized_mobile, jurisdiction_aeo.normalized_mobile, '') as aeo_mobile
                    from dormant_records dr
                    left join schools s on s.id = dr.school_ref
                    left join tehsils t on t.id = s.tehsil_ref
                    left join markazes m on m.id = s.markaz_ref
                    left join school_ddeo_overrides sdo
                      on sdo.school_ref = s.id and sdo.active = 1
                    left join ddeo_officers override_ddeo
                      on override_ddeo.id = sdo.ddeo_ref and override_ddeo.active = 1
                    left join ddeo_jurisdictions dj
                      on dj.wing_ref = s.wing_ref
                     and dj.tehsil_ref = s.tehsil_ref
                     and dj.active = 1
                     and sdo.id is null
                    left join ddeo_officers jurisdiction_ddeo
                      on jurisdiction_ddeo.id = dj.ddeo_ref and jurisdiction_ddeo.active = 1
                    left join school_aeo_overrides sao
                      on sao.school_ref = s.id and sao.active = 1
                    left join aeo_officers override_aeo
                      on override_aeo.id = sao.aeo_ref and override_aeo.active = 1
                    left join aeo_jurisdictions aj
                      on aj.wing_ref = s.wing_ref
                     and aj.markaz_ref = s.markaz_ref
                     and (aj.tehsil_ref = '' or aj.tehsil_ref = s.tehsil_ref)
                     and aj.active = 1
                     and sao.id is null
                    left join aeo_officers jurisdiction_aeo
                      on jurisdiction_aeo.id = aj.aeo_ref and jurisdiction_aeo.active = 1
                    where dr.run_key = ?
                    order by t.name, m.name, s.name
                    """,
                    (run_key,),
                ).fetchall()
            ]
    except sqlite3.Error as exc:
        return {}, f"Could not load dispatch context: {exc}"

    summary = _safe_json_loads(row["summary"])
    return {
        "run_key": run_key,
        "raw_file_name": str(row["raw_file_name"] or ""),
        "summary": summary,
        "rows": dormant_rows,
    }, ""
